import hashlib
import importlib.util
import io
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from cinesis.ranking import rank_loads
from cinesis.report import submission_note, write_outputs


def test_workbook_answers_and_preservation(
    tmp_path, workbook, profile, document, inputs
):
    before = hashlib.sha256(workbook.read_bytes()).hexdigest()
    result = rank_loads(profile, inputs[1], 15000)
    write_outputs(workbook, tmp_path, profile, document, result, None)
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == before
    with (
        zipfile.ZipFile(workbook) as source,
        zipfile.ZipFile(tmp_path / "completed.xlsx") as output,
    ):
        assert source.namelist() == output.namelist()
        changed = [
            name for name in source.namelist() if source.read(name) != output.read(name)
        ]
        assert len(changed) == 2
        assert all("worksheets/sheet" in name for name in changed)
    wb = load_workbook(tmp_path / "completed.xlsx", read_only=True, data_only=True)
    assert wb["Part A (Fill In)"]["B5"].value == "Dallas"
    assert wb["Part A (Fill In)"]["B13"].value.startswith("Unknown")
    assert wb["Part A (Fill In)"]["B16"].value == 15000
    assert wb["Part B (Fill In)"]["B5"].value == "L03"
    assert wb["Part B (Fill In)"]["C5"].value == "3.098"
    assert "CONDITIONAL" in wb["Part B (Fill In)"]["A2"].value
    assert "test fixture" in wb["Part B (Fill In)"]["A10"].value
    wb.close()


def test_strict_output_no_top_three(tmp_path, workbook, profile, document, inputs):
    result = rank_loads(profile, inputs[1])
    write_outputs(workbook, tmp_path, profile, document, result, None)
    wb = load_workbook(tmp_path / "completed.xlsx", read_only=True)
    assert wb["Part B (Fill In)"]["B5"].value == "Not established"
    assert wb["Part B (Fill In)"]["C5"].value is None
    wb.close()


def test_submission_and_readme_word_limit(profile, inputs):
    for capacity in [None, 15000]:
        note = submission_note(
            rank_loads(profile, inputs[1], capacity),
            "https://github.com/example/project",
            "openai",
            profile,
        )
        assert len(note.split()) <= 200
    assert (
        len((Path(__file__).resolve().parents[1] / "README.md").read_text().split())
        <= 200
    )


def test_secret_scan_including_xlsx():
    path = Path(__file__).resolve().parents[1] / "scripts/check_secrets.py"
    spec = importlib.util.spec_from_file_location("scanner", path)
    scanner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(scanner)
    fake = b"sk-" + b"x" * 32
    assert scanner.contains_secret(fake)
    assert not scanner.contains_secret(b"OPENAI_API_KEY=")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/worksheets/sheet1.xml", fake)
    assert scanner.contains_secret(buf.getvalue())
