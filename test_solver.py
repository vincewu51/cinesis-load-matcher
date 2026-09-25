import hashlib
import importlib.util
import io
import json
import math
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openai import OpenAIError
from openpyxl import load_workbook
from pydantic import ValidationError

import solve as extraction
from solve import DriverProfile, Evidence, haversine, rank_loads, validate_evidence
from workbook import city_coordinates, read_inputs, submission_note, transcript_hash, write_outputs

ROOT = Path(__file__).resolve().parent


@pytest.fixture
def workbook():
    return ROOT / "data/input.xlsx"


@pytest.fixture
def inputs(workbook):
    return read_inputs(workbook)


@pytest.fixture
def profile():
    return DriverProfile.model_validate(
        json.loads((ROOT / "results/profile.json").read_text())["profile"]
    )


@pytest.fixture
def document(profile, inputs):
    return {
        "schema_version": 1,
        "transcript_sha256": transcript_hash(inputs[0]),
        "extraction": {"kind": "test_fixture", "model": "none"},
        "profile": profile.model_dump(),
    }


def test_haversine_known_geometry():
    assert haversine((0, 0), (0, 0)) == 0
    assert haversine((0, 0), (0, 90)) == pytest.approx(6218.4071, abs=0.001)
    assert haversine((29, -98), (32, -96)) == haversine((32, -96), (29, -98))


def test_conditional_top_three_and_empty_return(profile, inputs):
    result = rank_loads(profile, inputs[1], 15000)
    assert result["mode"] == "conditional"
    assert [r["load_id"] for r in result["top_three"]] == ["L03", "L08", "L02"]
    assert [f"{r['effective_rate_per_mile']:.3f}" for r in result["top_three"]] == [
        "3.098",
        "2.480",
        "2.418",
    ]
    l08 = next(r for r in result["audit"] if r["load_id"] == "L08")
    assert l08["deadhead_to_origin_miles"] == 0
    assert l08["deadhead_home_miles"] == pytest.approx(223.111, abs=0.001)
    assert profile.weight_capacity_lb.value is None


def test_unknown_capacity_is_not_eligibility(profile, inputs):
    result = rank_loads(profile, inputs[1])
    assert result["mode"] == "needs_capacity"
    assert result["top_three"] == []
    assert [r["load_id"] for r in result["audit"] if r["status"] == "needs_capacity"] == [
        "L02",
        "L03",
        "L08",
    ]


@pytest.mark.parametrize(
    "capacity,expected",
    [
        (12000, ["L02"]),
        (12600, ["L08", "L02"]),
        (13000, ["L08", "L02"]),
        (14200, ["L03", "L08", "L02"]),
    ],
)
def test_capacity_sensitivity(profile, inputs, capacity, expected):
    assert [r["load_id"] for r in rank_loads(profile, inputs[1], capacity)["top_three"]] == expected


def test_equipment_missing_data_and_rejections(profile, inputs):
    rows = {r["load_id"]: r for r in rank_loads(profile, inputs[1], 15000)["audit"]}
    assert "incompatible equipment" in rows["L05"]["reasons"]
    assert "incompatible equipment" in rows["L04"]["reasons"]
    assert "effective rate does not meet" in " ".join(rows["L04"]["reasons"])
    assert "price" in " ".join(rows["L06"]["reasons"])
    assert "Destination" in " ".join(rows["L07"]["reasons"])


def test_strict_rate_boundary_uses_full_precision(profile, inputs):
    rows = inputs[1]
    rate = next(r for r in rank_loads(profile, rows, 15000)["audit"] if r["load_id"] == "L03")[
        "effective_rate_per_mile"
    ]
    profile.minimum_rate.value.dollars_per_mile = rate
    assert not rank_loads(profile, rows, 15000)["top_three"]
    profile.minimum_rate.value.comparison = ">="
    assert rank_loads(profile, rows, 15000)["top_three"][0]["load_id"] == "L03"


def test_confirmed_capacity_cannot_be_overridden(profile, inputs):
    profile.weight_capacity_lb.value = 13000
    profile.weight_capacity_lb.evidence = [Evidence(row=20, quote="example for unit test")]
    result = rank_loads(profile, inputs[1])
    assert result["mode"] == "verified_capacity"
    assert [r["load_id"] for r in result["top_three"]] == ["L08", "L02"]
    with pytest.raises(ValueError):
        rank_loads(profile, inputs[1], 15000)


@pytest.mark.parametrize("bad", [0, -1, math.nan, math.inf])
def test_invalid_capacity(profile, inputs, bad):
    with pytest.raises(ValueError):
        rank_loads(profile, inputs[1], bad)


@pytest.mark.parametrize(
    "field,value",
    [
        ("Price ($)", math.nan),
        ("Weight", -1),
        ("Dest Lat", 91),
        ("Dest Lon", -181),
        ("Destination", ""),
    ],
)
def test_malformed_load_excluded(profile, inputs, field, value):
    rows = inputs[1]
    rows[2][field] = value
    assert "L03" not in [r["load_id"] for r in rank_loads(profile, rows, 15000)["top_three"]]


def test_duplicate_ids_and_ties(profile, inputs):
    rows = inputs[1]
    copy = {**rows[2], "Load ID": "L00"}
    assert [r["load_id"] for r in rank_loads(profile, rows + [copy], 15000)["top_three"]][:2] == [
        "L00",
        "L03",
    ]
    duplicated = rank_loads(profile, rows + [rows[2]], 15000)
    assert "L03" not in [r["load_id"] for r in duplicated["top_three"]]


def test_zero_total_distance_excluded(profile, inputs):
    profile.current_location.value = "San Antonio"
    rows = inputs[1] + [
        {
            "Load ID": "ZERO",
            "Origin": "San Antonio",
            "Origin Lat": 29.4241,
            "Origin Lon": -98.4936,
            "Destination": "San Antonio",
            "Dest Lat": 29.4241,
            "Dest Lon": -98.4936,
            "Trailer": "Hotshot",
            "Weight": 1000,
            "Price ($)": 5000,
        }
    ]
    result = rank_loads(profile, rows, 15000)
    assert result["audit"][-1]["reasons"] == ["zero total trip distance"]


def test_coordinate_lookup_rejects_conflicts(inputs):
    rows = inputs[1]
    assert city_coordinates("Dallas, TX", rows) == (32.7767, -96.797)
    rows.append({**rows[-1], "Origin Lat": 33})
    with pytest.raises(ValueError):
        city_coordinates("Dallas", rows)


def mock_client(monkeypatch, response):
    monkeypatch.setenv("OPENAI_API_KEY", "unit-test-not-a-real-key")
    monkeypatch.setattr(extraction, "load_dotenv", lambda *a, **kw: None)
    client = MagicMock()
    client.__enter__.return_value = client
    client.responses.parse.return_value = response
    factory = MagicMock(return_value=client)
    monkeypatch.setattr(extraction, "OpenAI", factory)
    return client, factory


def test_real_transcript_evidence(profile, inputs):
    validate_evidence(profile, inputs[0])
    assert profile.weight_capacity_lb.value is None


def test_fabricated_quote_rejected(profile, inputs):
    profile.current_location.evidence = [Evidence(row=4, quote="I am in Boston")]
    with pytest.raises(ValueError, match="Unsupported evidence"):
        validate_evidence(profile, inputs[0])


def test_claim_requires_evidence(profile):
    data = profile.model_dump()
    data["weight_capacity_lb"]["value"] = 44000
    data["weight_capacity_lb"]["evidence"] = []
    with pytest.raises(ValidationError):
        DriverProfile.model_validate(data)


def test_llm_request_and_saved_provenance(monkeypatch, profile, inputs):
    client, factory = mock_client(
        monkeypatch,
        SimpleNamespace(output_parsed=profile, status="completed", model="test-model"),
    )
    document = extraction.extract(inputs[0], "test-model")
    args = client.responses.parse.call_args.kwargs
    assert args["store"] is False
    assert args["text_format"] is DriverProfile
    assert json.loads(args["input"][1]["content"]) == inputs[0]
    assert "unit-test-not-a-real-key" not in json.dumps(document)
    assert document["extraction"]["kind"] == "openai"
    assert document["profile"]["weight_capacity_lb"]["value"] is None
    assert factory.call_args.kwargs["base_url"] == "https://api.openai.com/v1"


def test_changed_transcript_not_hardcoded(monkeypatch, profile, inputs):
    conversation = inputs[0]
    conversation[2]["dialogue"] = conversation[2]["dialogue"].replace("Dallas", "Houston")
    profile.current_location.value = "Houston"
    profile.current_location.evidence = [Evidence(row=4, quote="I'm in Houston.")]
    mock_client(
        monkeypatch,
        SimpleNamespace(output_parsed=profile, status="completed", model="test-model"),
    )
    result = extraction.extract(conversation, "test-model")
    assert result["profile"]["current_location"]["value"] == "Houston"


@pytest.mark.parametrize("status,parsed", [("incomplete", True), ("completed", False)])
def test_incomplete_or_refused_response(monkeypatch, profile, inputs, status, parsed):
    mock_client(
        monkeypatch,
        SimpleNamespace(
            output_parsed=profile if parsed else None, status=status, model="test-model"
        ),
    )
    with pytest.raises(extraction.ExtractionError, match="incomplete"):
        extraction.extract(inputs[0], "test-model")


def test_provider_errors_do_not_expose_payloads(monkeypatch, profile, inputs):
    client, _ = mock_client(monkeypatch, None)
    client.responses.parse.side_effect = OpenAIError("sensitive-marker")
    with pytest.raises(extraction.ExtractionError) as exc:
        extraction.extract(inputs[0], "test-model")
    assert "sensitive-marker" not in str(exc.value)


def test_missing_key_no_network(monkeypatch, inputs):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.setattr(extraction, "load_dotenv", lambda *a, **kw: None)
    client = MagicMock()
    monkeypatch.setattr(extraction, "OpenAI", client)
    with pytest.raises(extraction.ExtractionError, match="Configure"):
        extraction.extract(inputs[0])
    client.assert_not_called()


def test_stale_saved_profile_rejected(tmp_path, document, inputs):
    path = tmp_path / "profile.json"
    path.write_text(json.dumps(document))
    inputs[0][0]["dialogue"] = "changed"
    with pytest.raises(ValueError, match="fingerprint"):
        extraction.read_profile(path, inputs[0])


def test_workbook_answers_and_preservation(tmp_path, workbook, profile, document, inputs):
    before = hashlib.sha256(workbook.read_bytes()).hexdigest()
    result = rank_loads(profile, inputs[1], 15000)
    write_outputs(workbook, tmp_path, profile, document, result, None)
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == before
    with (
        zipfile.ZipFile(workbook) as source,
        zipfile.ZipFile(tmp_path / "completed.xlsx") as output,
    ):
        assert source.namelist() == output.namelist()
        changed = [name for name in source.namelist() if source.read(name) != output.read(name)]
        assert len(changed) == 2
        assert all("worksheets/sheet" in name for name in changed)
    wb = load_workbook(tmp_path / "completed.xlsx", read_only=True, data_only=True)
    assert wb["Part A (Fill In)"]["B5"].value == "Dallas"
    assert wb["Part A (Fill In)"]["B13"].value.startswith("Unknown")
    assert wb["Part A (Fill In)"]["B16"].value == 15000
    assert wb["Part B (Fill In)"]["B5"].value == "L03"
    assert wb["Part B (Fill In)"]["C5"].value == "3.098"
    assert "CONDITIONAL" in wb["Part B (Fill In)"]["A2"].value
    assert "test fixture" in wb["Part B (Fill In)"]["A10"].value
    wb.close()


def test_strict_output_no_top_three(tmp_path, workbook, profile, document, inputs):
    result = rank_loads(profile, inputs[1])
    write_outputs(workbook, tmp_path, profile, document, result, None)
    wb = load_workbook(tmp_path / "completed.xlsx", read_only=True)
    assert wb["Part B (Fill In)"]["B5"].value == "Not established"
    assert wb["Part B (Fill In)"]["C5"].value is None
    wb.close()


def test_submission_note_word_limit(profile, inputs):
    for capacity in [None, 15000]:
        note = submission_note(
            rank_loads(profile, inputs[1], capacity),
            "https://github.com/example/project",
            "openai",
            profile,
        )
        assert len(note.split()) <= 200


def test_secret_scan_including_xlsx():
    path = Path(__file__).resolve().parent / "security.py"
    spec = importlib.util.spec_from_file_location("scanner", path)
    scanner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(scanner)
    fake = b"sk-" + b"x" * 32
    assert scanner.contains_secret(fake)
    assert not scanner.contains_secret(b"OPENAI_API_KEY=")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/worksheets/sheet1.xml", fake)
    assert scanner.contains_secret(buf.getvalue())


def test_catalog_minimum_has_provenance_and_keeps_profile_unknown(profile, inputs):
    path = ROOT / "data/truck-capacities.json"
    result = rank_loads(profile, inputs[1], capacity_catalog=path)
    assert result["capacity_lb"] == 8710
    assert result["capacity_source"] == "catalog_minimum"
    assert result["mode"] == "conditional"
    assert profile.weight_capacity_lb.value is None
    estimate = result["capacity_estimate"]
    assert estimate["selected_record_id"] == "bigtex-14gn"
    assert estimate["catalog_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert len(estimate["matching_record_ids"]) == 5
    assert estimate["source_url"].startswith("https://www.bigtextrailers.com/")
    assert result["top_three"] == []
    rows = {r["load_id"]: r for r in result["audit"]}
    assert all(rows[key]["status"] == "needs_capacity" for key in ["L02", "L03", "L08"])
    assert rows["L04"]["status"] == "rejected"


def test_catalog_estimate_allows_only_conditional_candidates(profile, inputs):
    rows = inputs[1]
    rows[2]["Weight"] = 8000
    result = rank_loads(profile, rows, capacity_catalog=ROOT / "data/truck-capacities.json")
    assert result["top_three"][0]["load_id"] == "L03"
    assert result["top_three"][0]["status"] == "conditional"


@pytest.mark.parametrize("labels", [["Van"], ["Flatbed"], ["Reefer"], ["Gooseneck", "Van"], []])
def test_catalog_does_not_fill_uncovered_classes(labels):
    estimate = extraction.catalog_capacity(labels, ROOT / "data/truck-capacities.json")
    assert estimate["value_lb"] is None


def test_conflicting_catalog_source_blocks_box_truck_estimate():
    estimate = extraction.catalog_capacity(["Box Truck"], ROOT / "data/truck-capacities.json")
    assert estimate["value_lb"] is None
    assert "conflicts" in estimate["reason"]


def test_stated_capacity_precedes_catalog(profile, inputs):
    profile.weight_capacity_lb.value = 13000
    result = rank_loads(profile, inputs[1], capacity_catalog=Path("unused-path.json"))
    assert result["capacity_lb"] == 13000
    assert result["capacity_source"] == "transcript"
    assert result["capacity_estimate"] is None
    with pytest.raises(ValueError):
        rank_loads(profile, inputs[1], 15000, ROOT / "data/truck-capacities.json")


@pytest.mark.parametrize("bad", [-1, float("nan"), float("inf"), 20000])
def test_invalid_catalog_capacity_rejected(tmp_path, bad):
    catalog = json.loads((ROOT / "data/truck-capacities.json").read_text())
    catalog["records"][0]["payload_lb_min"] = bad
    path = tmp_path / "catalog.json"
    path.write_text(json.dumps(catalog))
    with pytest.raises(ValueError):
        extraction.catalog_capacity(["Gooseneck"], path)


def test_catalog_workbook_discloses_estimate(tmp_path, workbook, profile, document, inputs):
    result = rank_loads(profile, inputs[1], capacity_catalog=ROOT / "data/truck-capacities.json")
    write_outputs(workbook, tmp_path, profile, document, result)
    wb = load_workbook(tmp_path / "completed.xlsx", read_only=True, data_only=True)
    assert wb["Part A (Fill In)"]["B13"].value == "Unknown — not stated"
    assert wb["Part A (Fill In)"]["A16"].value == "Catalog estimate (not confirmed)"
    assert wb["Part A (Fill In)"]["B16"].value == 8710
    note = wb["Part B (Fill In)"]["A10"].value
    assert "not a guaranteed truck limit" in note
    assert "bigtextrailers.com" in note
    assert len(note.split()) <= 200
    wb.close()
