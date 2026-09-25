"""Read task inputs; populate answer cells without rewriting unrelated Excel content."""

import hashlib
import json
import math
import posixpath
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

HEADERS = [
    "Load ID",
    "Origin",
    "Origin Lat",
    "Origin Lon",
    "Destination",
    "Dest Lat",
    "Dest Lon",
    "Trailer",
    "Weight",
    "Price ($)",
]


def read_inputs(path: Path) -> tuple[list[dict], list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        conversation = [
            {"row": i, "speaker": str(row[0]), "dialogue": str(row[1])}
            for i, row in enumerate(
                workbook["Sample Conversation"].iter_rows(max_col=2, values_only=True),
                1,
            )
            if i > 1 and row[0] and row[1]
        ]
        rows = list(workbook["Loads"].iter_rows(max_col=10, values_only=True))
        if list(rows[0]) != HEADERS:
            raise ValueError("Unexpected Loads column headers")
        loads = [
            dict(zip(HEADERS, row))
            for row in rows[1:]
            if any(v is not None for v in row)
        ]
        if not conversation or not loads:
            raise ValueError("Workbook has no conversation or loads")
        return conversation, loads
    finally:
        workbook.close()


def transcript_hash(conversation: list[dict]) -> str:
    return hashlib.sha256(
        json.dumps(conversation, ensure_ascii=False, sort_keys=True).encode()
    ).hexdigest()


def normalize_city(name: str) -> str:
    return (
        " ".join(name.lower().replace(",", " ").split())
        .removesuffix(" texas")
        .removesuffix(" tx")
    )


def city_coordinates(city: str, loads: list[dict]) -> tuple[float, float]:
    matches = set()
    for load in loads:
        for field, lat, lon in [
            ("Origin", "Origin Lat", "Origin Lon"),
            ("Destination", "Dest Lat", "Dest Lon"),
        ]:
            if normalize_city(str(load[field])) == normalize_city(city):
                try:
                    point = (float(load[lat]), float(load[lon]))
                except (ValueError, TypeError):
                    continue
                if (
                    all(math.isfinite(v) for v in point)
                    and -90 <= point[0] <= 90
                    and -180 <= point[1] <= 180
                ):
                    matches.add(point)
    if len(matches) != 1:
        raise ValueError(f"Expected one unambiguous coordinate pair for {city}")
    return next(iter(matches))


def fill_workbook(source: Path, target: Path, answers: dict[str, dict]) -> None:
    """Patch only values in the two answer sheets; preserve drawings, styles, and all other ZIP parts."""
    if source.resolve() == target.resolve():
        raise ValueError("Output must not overwrite the original workbook")
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    namespace = ns["m"]
    with zipfile.ZipFile(source) as archive:
        relationships = {
            r.get("Id"): r.get("Target")
            for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        }
        replacements = {}
        for sheet in ET.fromstring(archive.read("xl/workbook.xml")).find(
            "m:sheets", ns
        ):
            if sheet.get("name") not in answers:
                continue
            rel = sheet.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            )
            name = relationships[rel]
            name = (
                name.lstrip("/")
                if name.startswith("/")
                else posixpath.normpath("xl/" + name)
            )
            root = ET.fromstring(archive.read(name))
            data = root.find("m:sheetData", ns)
            for address, value in answers[sheet.get("name")].items():
                row_number = int("".join(c for c in address if c.isdigit()))
                row = next((r for r in data if int(r.get("r")) == row_number), None)
                if row is None:
                    row = ET.Element(f"{{{namespace}}}row", {"r": str(row_number)})
                    index = next(
                        (i for i, r in enumerate(data) if int(r.get("r")) > row_number),
                        len(data),
                    )
                    data.insert(index, row)
                cell = next((c for c in row if c.get("r") == address), None)
                if cell is None:
                    cell = ET.SubElement(row, f"{{{namespace}}}c", {"r": address})
                for child in list(cell):
                    cell.remove(child)
                cell.attrib.pop("t", None)
                if value is None:
                    continue
                if isinstance(value, (int, float)):
                    ET.SubElement(cell, f"{{{namespace}}}v").text = str(value)
                else:
                    cell.set("t", "inlineStr")
                    text = ET.SubElement(
                        ET.SubElement(cell, f"{{{namespace}}}is"), f"{{{namespace}}}t"
                    )
                    text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    text.text = str(value)
            replacements[name] = ET.tostring(
                root, encoding="utf-8", xml_declaration=True
            )
        if len(replacements) != len(answers):
            raise ValueError("Missing expected answer sheet")
        with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as output:
            for entry in archive.infolist():
                output.writestr(
                    entry,
                    replacements.get(entry.filename, archive.read(entry.filename)),
                )
